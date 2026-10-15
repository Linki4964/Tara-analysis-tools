from html import escape
from pathlib import Path
from typing import Any
from zipfile import ZipFile
import csv
import io
import json
import re
import xml.etree.ElementTree as ET

from fastapi import HTTPException, UploadFile


WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


async def extract_upload(file: UploadFile) -> dict[str, Any]:
    filename = file.filename or "uploaded"
    ext = Path(filename).suffix.lower()

    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File is too large. Maximum size is 50MB.")

    extracted_text = ""
    extracted_html = ""

    try:
        if ext == ".docx":
            extracted = _extract_docx(content)
            extracted_text = extracted["text"]
            extracted_html = extracted["html"]
        elif ext == ".pdf":
            extracted_text = _extract_pdf(content)
        elif ext in {".xlsx", ".xls"}:
            extracted_text = _extract_excel(content, ext)
        else:
            extracted_text = _extract_generic_text(content, filename)
    except Exception:
        # RAG ingestion is intentionally broad: a malformed or uncommon file
        # still gets indexed through the best-effort generic extractor.
        extracted_text = _extract_generic_text(content, filename)

    extracted_text = extracted_text.strip()
    if len(extracted_text) < 10:
        extracted_text = f"文件：{filename}\n文件类型：{ext or '未知'}\n未提取到可读文本。"

    return {
        "success": True,
        "metadata": {
            "filename": filename,
            "fileType": ext,
            "fileSize": len(content),
            "charCount": len(extracted_text),
            "lineCount": len(extracted_text.splitlines()),
            "hasHtml": bool(extracted_html),
        },
        "extractedText": extracted_text,
        "extractedHtml": extracted_html or None,
}


def _extract_generic_text(content: bytes, filename: str) -> str:
    """Best-effort text extraction for any extension.

    Plain text, HTML, XML, source code, logs and CSV-like files are decoded
    directly.  For binary formats without a dedicated parser, printable text
    fragments are recovered so the upload can still be indexed instead of
    failing solely because of its filename extension.
    """
    for encoding in ("utf-8", "utf-16", "gb18030", "big5"):
        try:
            decoded = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        printable = sum(char.isprintable() or char in "\n\r\t" for char in decoded)
        if decoded and printable / len(decoded) >= 0.72:
            return decoded

    fragments = re.findall(rb"[\x20-\x7e\x80-\xff]{4,}", content)
    recovered = []
    for fragment in fragments[:2000]:
        for encoding in ("utf-8", "gb18030", "latin-1"):
            try:
                value = fragment.decode(encoding).strip()
                if value:
                    recovered.append(value)
                break
            except UnicodeDecodeError:
                continue
    if recovered:
        return "\n".join(recovered)
    return f"文件：{filename}\n该文件为二进制或图像格式，未发现可提取的文本片段。"


def _extract_docx(content: bytes) -> dict[str, str]:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            document_xml = archive.read("word/document.xml")
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Failed to parse DOCX file: {error}") from error

    root = ET.fromstring(document_xml)
    paragraphs: list[str] = []
    html_parts: list[str] = []

    for child in root.findall(".//w:body/*", WORD_NS):
        tag = _local_name(child.tag)
        if tag == "p":
            text = _node_text(child)
            if text:
                paragraphs.append(text)
                html_parts.append(f"<p>{escape(text)}</p>")
        elif tag == "tbl":
            rows = []
            for row in child.findall(".//w:tr", WORD_NS):
                cells = [_node_text(cell) for cell in row.findall("./w:tc", WORD_NS)]
                rows.append(cells)
                if any(cells):
                    paragraphs.append(" | ".join(cells))
            if rows:
                html_parts.append(_table_to_html(rows))

    return {"text": "\n".join(paragraphs), "html": "\n".join(html_parts)}


def _node_text(node: ET.Element) -> str:
    pieces = [text_node.text or "" for text_node in node.findall(".//w:t", WORD_NS)]
    return "".join(pieces).strip()


def _table_to_html(rows: list[list[str]]) -> str:
    row_html = []
    for row in rows:
        cells = "".join(f"<td>{escape(cell)}</td>" for cell in row)
        row_html.append(f"<tr>{cells}</tr>")
    return f"<table>{''.join(row_html)}</table>"


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise HTTPException(
            status_code=500,
            detail="PDF extraction requires pypdf. Please install backend dependencies from requirements.txt.",
        ) from error
    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail="Failed to parse PDF. The file may be encrypted, scanned, or corrupted.",
        ) from error


def _extract_excel(content: bytes, ext: str) -> str:
    if ext == ".xls":
        raise HTTPException(status_code=400, detail="旧版 .xls 暂不支持，请另存为 .xlsx 或 CSV 后上传。")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    lines.append(" | ".join(values))
        return "\n".join(lines)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {error}") from error
